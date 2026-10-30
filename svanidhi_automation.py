#!/usr/bin/env python3
"""
PM SVANidhi - Vendor Scheme Tracker automation
==============================================

What it does (mirrors your screen recording exactly)
----------------------------------------------------
For every Application No. in data.xlsx that has rows with Status = "Eligible":
  1. Search the Application No. on the portal's Vendor Scheme Tracker.
  2. For every Eligible scheme row: click Apply -> enter Reference No -> Submit.
     (verifies the stage becomes "Applied")
  3. For every Applied scheme row: click the thumbs-up -> confirm "Yes".
     (verifies the stage becomes "Sanctioned")
  4. Write "Sanction" into the Status column of that Excel row (only after the
     portal confirmed it), and save.
It works for any number of schemes per application and is safe to re-run:
rows already Applied / Sanctioned on the portal are resumed or skipped.

One-time setup
--------------
  pip install playwright openpyxl          (no "playwright install" needed)

Every run
---------
  1. Close data.xlsx in Excel (the script needs to save into it).
  2. Start Chrome with a debugging port (Windows, run in Command Prompt or 1_START_CHROME.bat):
       "C:\\Program Files\\Google\\Chrome\\Application\\chrome.exe" ^
         --remote-debugging-port=9222 --user-data-dir="C:\\chrome-svanidhi"
  3. In THAT Chrome window log in to the portal yourself (captcha/OTP) and open
     "Vendor Scheme Tracker".
  4. First do a dry run (searches and reads only, no changes):
       python svanidhi_automation.py --excel data.xlsx.xlsx --dry-run
  5. Then a small real run, then the full run:
       python svanidhi_automation.py --excel data.xlsx.xlsx --limit 2
       python svanidhi_automation.py --excel data.xlsx.xlsx

If a selector does not match your portal, run:
       python svanidhi_automation.py --excel data.xlsx.xlsx --inspect PMS072615318872
and send me the printed output; all selectors are in the SEL block below.

Credentials are never read from the Excel sheet or stored by this script.
"""
from __future__ import annotations

import argparse
import csv
import re
import shutil
import sys
import time
from collections import Counter, OrderedDict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

# =============================== CONFIG ======================================

# Excel header names (lower-case, trailing dots ignored). Add aliases if needed.
EXCEL_HEADERS = {
    "vendor": ["vendor name"],
    "scheme": ["scheme name"],
    "app": ["application no", "application number"],
    "ref": ["reference no", "reference number", "reference"],
    "status": ["status"],
}
PENDING_STATUS = "eligible"      # Excel Status that means "still to do"
DONE_STATUS_TEXT = "Sanction"    # what gets written back after success

# If Excel and portal spell a scheme differently, map them here (lower-case):
#   {"pm jivan jyoti bima yojana": "pm jeevan jyoti bima yojana"}
SCHEME_ALIASES: dict[str, str] = {
    "pm suraksha bima yojana for mother": "pm suraksha bima for mother",
    "pm suraksha bima yojana for father": "pm suraksha bima for father",
}

# Sanity check that the Reference No belongs to the scheme (catches shifted rows).
# (keyword in scheme name, required prefix of reference no). Disable with
# --no-prefix-check, or edit for other sheets (Labour, Health, WCD, Food).
REF_PREFIX_RULES = [
    ("suraksha bima", "PMSBY"),
    ("jeevan jyoti", "PMJJBY"),
    ("jyoti", "PMJJBY"),
]

# Portal selectors - the only place to change if the page differs.
SEL = {
    "table": 'table:has(th:has-text("Scheme Stage"))',
    "app_input": ['input[placeholder*="Application" i]', 'input[name*="application" i]'],
    "search_btn": 'button:has-text("Search"), input[type="submit"][value*="Search" i]',
    "icon": "a, button, i",                       # clickable icon inside a cell
    "dialog": '.modal.show, .modal.in, [role="dialog"], .swal2-popup, .bootbox',
    "dialog_input": 'input[type="text"], input:not([type]), textarea',
    "dialog_button": "button, a, input[type=submit]",
}
APPLY_DIALOG_TEXT = r"Update Reference No"
SANCTION_DIALOG_TEXT = r"mark the selected application as"

# ============================== HELPERS ======================================


class PortalError(Exception):
    """Something on the portal did not behave as expected (app is skipped)."""


def norm(s) -> str:
    return re.sub(r"\s+", " ", str(s if s is not None else "").replace("\xa0", " ")).strip().lower()


def canon(s) -> str:
    n = norm(s)
    return SCHEME_ALIASES.get(n, n)


def appkey(s) -> str:
    return re.sub(r"\s+", "", str(s if s is not None else "")).upper()


@dataclass
class Item:
    row: int
    vendor: str
    scheme: str       # canonical (lower-case) scheme name
    scheme_raw: str
    app: str
    ref: str
    status: str


# ================================ EXCEL ======================================


def load_excel(path: Path, sheet: str):
    # Auto-detect if user passed data.xlsx but actual file on Windows is data.xlsx.xlsx
    if not path.exists():
        double_ext = path.with_name(path.name + ".xlsx")
        if double_ext.exists():
            print(f"Notice: '{path.name}' was not found, using '{double_ext.name}' instead.")
            path = double_ext

    wb = load_workbook(path)
    if sheet not in wb.sheetnames:
        # Check if double-ext file has the sheet
        double_ext = path.with_name(path.name + ".xlsx")
        if double_ext.exists() and double_ext != path:
            wb_alt = load_workbook(double_ext)
            if sheet in wb_alt.sheetnames:
                print(f"Notice: Sheet '{sheet}' found in '{double_ext.name}', switching to '{double_ext.name}'.")
                wb = wb_alt
                path = double_ext

    if sheet not in wb.sheetnames:
        sys.exit(f"Sheet '{sheet}' not found in '{path.name}'. Sheets available: {wb.sheetnames}")
    ws = wb[sheet]
    hdr_row, cols = None, {}
    for r in range(1, 8):
        names = [norm(c.value).rstrip(".") for c in ws[r]]
        found = {}
        for key, aliases in EXCEL_HEADERS.items():
            for idx, name in enumerate(names, start=1):
                if name in [a.rstrip(".") for a in aliases]:
                    found[key] = idx
                    break
        if len(found) == len(EXCEL_HEADERS):
            hdr_row, cols = r, found
            break
    if not hdr_row:
        sys.exit(f"Could not find header row with columns: {list(EXCEL_HEADERS)} in sheet '{sheet}' of '{path.name}'")

    items = []
    for r in range(hdr_row + 1, ws.max_row + 1):
        app = appkey(ws.cell(r, cols["app"]).value)
        if not app:
            continue
        sch_raw = str(ws.cell(r, cols["scheme"]).value or "").strip()
        ref_val = ws.cell(r, cols["ref"]).value
        if ref_val is None:
            ref_str = ""
        elif isinstance(ref_val, float) and ref_val.is_integer():
            ref_str = str(int(ref_val))
        else:
            ref_str = re.sub(r"\s+", "", str(ref_val))

        items.append(Item(
            row=r,
            vendor=str(ws.cell(r, cols["vendor"]).value or "").strip(),
            scheme=canon(sch_raw),
            scheme_raw=sch_raw,
            app=app,
            ref=ref_str,
            status=norm(ws.cell(r, cols["status"]).value),
        ))
    return wb, ws, cols, items, path


def check_ref(item: Item, enabled: bool):
    if not item.ref:
        raise PortalError(f"row {item.row}: Reference No is empty")
    if not enabled:
        return
    for kw, prefix in REF_PREFIX_RULES:
        if kw in item.scheme and not item.ref.upper().startswith(prefix):
            raise PortalError(
                f"row {item.row}: Reference No '{item.ref}' does not start with {prefix} "
                f"but scheme is '{item.scheme_raw}' (Excel rows misaligned?)")


# ================================ PORTAL =====================================


@dataclass
class Row:
    i: int
    scheme: str
    app: str
    stage: str


def app_input(page):
    for s in SEL["app_input"]:
        loc = page.locator(s).first
        if loc.count() and loc.is_visible():
            return loc
    # Fallback: filter row order is Vendor Name | Application No | Mobile No
    inputs = page.locator('input[type="text"], input:not([type])').locator("visible=true")
    if inputs.count() >= 3:
        return inputs.nth(1)
    raise PortalError("Application No. search box not found (are you on Vendor Scheme Tracker and logged in?)")


def column_map(heads):
    m = {}
    for i, h in enumerate(heads):
        if h == "scheme":
            m["scheme"] = i
        elif "application" in h:
            m["app"] = i
        elif "stage" in h:
            m["stage"] = i
        elif h == "apply":
            m["apply"] = i
        elif "sanction" in h:
            m["sanction"] = i
    missing = [k for k in ("scheme", "app", "stage", "apply", "sanction") if k not in m]
    if missing:
        raise PortalError(f"Table columns not found: {missing}; headers seen: {heads}")
    return m


def disable_animations(page):
    """Disables CSS modal fade/transition animations so popups open and close in 0ms without affecting accuracy."""
    try:
        page.evaluate("""() => {
            if (!document.getElementById('svanidhi-fast-style')) {
                const style = document.createElement('style');
                style.id = 'svanidhi-fast-style';
                style.textContent = '*, *::before, *::after { transition-duration: 0s !important; animation-duration: 0s !important; }';
                document.head.appendChild(style);
            }
        }""")
    except Exception:
        pass


def table(page):
    tbl = page.locator(SEL["table"]).first
    tbl.wait_for(state="visible", timeout=10000)
    return tbl


def read_table(page):
    data = page.evaluate("""() => {
        const tbl = Array.from(document.querySelectorAll('table')).find(t =>
            /scheme\\s+stage/i.test(t.innerText || '')
        );
        if (!tbl) return null;
        const heads = Array.from(tbl.querySelectorAll('thead th')).map(th => th.innerText || '');
        const rows = Array.from(tbl.querySelectorAll('tbody tr')).map(tr =>
            Array.from(tr.querySelectorAll('td')).map(td => td.innerText || '')
        );
        return { heads, rows };
    }""")
    if not data:
        raise PortalError("Scheme table not visible yet")
    heads = [norm(h) for h in data["heads"]]
    col = column_map(heads)
    max_col = max(col.values())
    rows = []
    for i, tds in enumerate(data["rows"]):
        if len(tds) <= max_col:
            continue  # "No data" row
        rows.append(Row(i, canon(tds[col["scheme"]]), appkey(tds[col["app"]]), norm(tds[col["stage"]])))
    return rows, col


def footer_total(page):
    try:
        t = page.locator(r"text=/Showing\s+\d+\s+to\s+\d+\s+of\s+\d+\s+entries/i").first.inner_text(timeout=2000)
        return int(re.search(r"of\s+(\d+)", t).group(1))
    except Exception:
        return None


def set_max_page_size(page):
    try:
        sel = page.locator("select").filter(has=page.locator("option", has_text=re.compile("show entries", re.I))).first
        opts = sel.locator("option").evaluate_all("els=>els.map(e=>[e.value,e.textContent.trim()])")
        nums = [(int(v), v) for v, t in opts if re.fullmatch(r"\d+", v or "")]
        if nums:
            sel.select_option(value=max(nums)[1])
            page.wait_for_timeout(1200)
    except Exception as e:  # not fatal; the footer check below still guards us
        print(f"   (could not change 'Show Entries': {e})")


def search(page, app: str, delay: float):
    inp = app_input(page)
    inp.fill(app)
    page.locator(SEL["search_btn"]).first.click()
    deadline = time.time() + 20
    rows = col = None
    while time.time() < deadline:
        time.sleep(0.05)
        try:
            rows, col = read_table(page)
        except Exception:
            continue
        # results must belong to the application we just searched (not the previous one)
        if rows and all(r.app == app for r in rows):
            break
    else:
        raise PortalError(f"Search results for {app} did not load / no rows returned")
    # Pagination check only needed if 10+ rows are visible (minimum portal page size is 10)
    if len(rows) >= 10:
        total = footer_total(page)
        if total is not None and total != len(rows):
            set_max_page_size(page)
            rows, col = read_table(page)
            total = footer_total(page)
            if total is not None and total != len(rows):
                raise PortalError(f"Portal lists {total} entries but only {len(rows)} rows are visible (pagination)")
    if delay > 0:
        time.sleep(delay)
    return rows, col


def pick_row(rows, app: str, scheme: str):
    hits = [r for r in rows if r.app == app and r.scheme == scheme]
    if len(hits) != 1:
        raise PortalError(f"Expected exactly 1 portal row for '{scheme}', found {len(hits)}")
    return hits[0]


def find_row(page, app: str, scheme: str):
    rows, col = read_table(page)
    return pick_row(rows, app, scheme), col


def wait_dialog(page, text_regex: str, timeout=8000):
    dlg = page.locator(SEL["dialog"]).filter(has_text=re.compile(text_regex, re.I)).first
    dlg.wait_for(state="visible", timeout=timeout)
    return dlg


def dialog_button(dlg, label_regex: str):
    return dlg.locator(SEL["dialog_button"]).filter(has_text=re.compile(label_regex, re.I)).first


def wait_stage(page, app: str, scheme: str, want: str, delay: float):
    """Poll rapidly until the scheme row shows the wanted stage (re-pressing Search quickly if stale)."""
    for attempt in range(2):
        end = time.time() + (3.5 if attempt == 0 else 10)
        while time.time() < end:
            time.sleep(0.05)
            try:
                rows, col = read_table(page)
                row = pick_row(rows, app, scheme)
            except Exception:
                continue
            if row.stage == want:
                if delay > 0:
                    time.sleep(delay)
                return rows, col
        page.locator(SEL["search_btn"]).first.click()
    raise PortalError(f"'{scheme}' did not reach stage '{want}'")


def do_apply(page, app: str, scheme: str, ref: str, delay: float, row=None, col=None):
    if row is None or col is None:
        row, col = find_row(page, app, scheme)
    cell = table(page).locator("tbody tr").nth(row.i).locator("td").nth(col["apply"])
    cell.locator(SEL["icon"]).first.click()
    dlg = wait_dialog(page, APPLY_DIALOG_TEXT)
    box = dlg.locator(SEL["dialog_input"]).first
    box.fill(ref)
    if box.input_value().strip() != ref:
        raise PortalError("Reference No was not typed correctly into the popup")
    dialog_button(dlg, r"^\s*submit\s*$").click()
    dlg.wait_for(state="hidden", timeout=10000)
    return wait_stage(page, app, scheme, "applied", delay)


def do_sanction(page, app: str, scheme: str, delay: float, row=None, col=None):
    if row is None or col is None:
        row, col = find_row(page, app, scheme)
    cell = table(page).locator("tbody tr").nth(row.i).locator("td").nth(col["sanction"])
    cell.locator(SEL["icon"]).first.click()
    dlg = wait_dialog(page, SANCTION_DIALOG_TEXT)
    text = norm(dlg.inner_text())
    if "sanctioned" not in text:  # never confirm a Reject / unknown dialog
        page.keyboard.press("Escape")
        raise PortalError(f"Unexpected confirmation dialog: '{text[:80]}'")
    dialog_button(dlg, r"^\s*yes\s*$").click()
    dlg.wait_for(state="hidden", timeout=10000)
    return wait_stage(page, app, scheme, "sanctioned", delay)


# ============================== MAIN LOGIC ===================================


class Logger:
    def __init__(self, logdir: Path):
        logdir.mkdir(parents=True, exist_ok=True)
        self.dir = logdir
        ts = datetime.now().strftime("%Y%m%d-%H%M%S")
        self.path = logdir / f"run-{ts}.csv"
        self.f = open(self.path, "w", newline="", encoding="utf-8")
        self.w = csv.writer(self.f)
        self.w.writerow(["time", "excel_row", "application_no", "scheme", "action", "result", "detail"])

    def __call__(self, row, app, scheme, action, result, detail=""):
        self.w.writerow([datetime.now().isoformat(timespec="seconds"), row, app, scheme, action, result, detail])
        self.f.flush()
        print(f"   [{result:7}] {action:10} {scheme}  {detail}")


def process_app(page, app, all_items, args, ws, cols, log):
    pending = [i for i in all_items if i.status == PENDING_STATUS]

    # ---- sanity checks on the Excel side
    dup = [s for s, n in Counter(i.scheme for i in all_items).items() if n > 1]
    if dup:
        raise PortalError(f"Same scheme appears more than once in Excel for this application: {dup}")
    for it in pending:
        check_ref(it, not args.no_prefix_check)

    # ---- search and reconcile with the portal
    rows, col = search(page, app, args.delay)
    portal = Counter(r.scheme for r in rows)
    dup_p = [s for s, n in portal.items() if n > 1]
    if dup_p:
        raise PortalError(f"Duplicate scheme rows on portal: {dup_p}")
    excel_names = {i.scheme for i in all_items}
    missing = [i.scheme_raw for i in pending if i.scheme not in portal]
    extra = [s for s in portal if s not in excel_names]
    if missing:
        raise PortalError(f"Scheme(s) in Excel but not on portal: {missing} (portal has: {list(portal)})")
    if extra and args.strict_portal:
        raise PortalError(f"Scheme(s) on portal but not in Excel: {extra}")
    elif extra:
        print(f"   [INFO   ] skip       Ignoring extra scheme(s) on portal not in Excel: {extra}")

    def stage_of(it):
        return next(r.stage for r in rows if r.scheme == it.scheme)

    # ---- dry run: just print the plan
    if args.dry_run:
        for it in pending:
            st = stage_of(it)
            plan = {"eligible": f"apply ref {it.ref} then sanction",
                    "applied": "sanction only", "sanctioned": "already sanctioned -> mark Excel"}.get(st, f"UNEXPECTED stage '{st}'")
            log(it.row, app, it.scheme_raw, "plan", "DRY", plan)
        return False

    # ---- phase 1: Apply (reference number entry)
    for it in pending:
        row = pick_row(rows, app, it.scheme)
        if row.stage == "eligible":
            rows, col = do_apply(page, app, it.scheme, it.ref, args.delay, row=row, col=col)
            log(it.row, app, it.scheme_raw, "apply", "OK", it.ref)
        elif row.stage in ("applied", "sanctioned"):
            log(it.row, app, it.scheme_raw, "apply", "SKIP", f"already {row.stage}")
        else:
            raise PortalError(f"'{it.scheme_raw}' has unexpected stage '{row.stage}'")

    # ---- phase 2: Sanction
    for it in pending:
        row = pick_row(rows, app, it.scheme)
        if row.stage == "applied":
            rows, col = do_sanction(page, app, it.scheme, args.delay, row=row, col=col)
            log(it.row, app, it.scheme_raw, "sanction", "OK")
        elif row.stage == "sanctioned":
            log(it.row, app, it.scheme_raw, "sanction", "SKIP", "already sanctioned")
        else:
            raise PortalError(f"'{it.scheme_raw}' is '{row.stage}', expected applied/sanctioned")
        # portal confirmed -> record in Excel memory
        ws.cell(it.row, cols["status"]).value = DONE_STATUS_TEXT
        log(it.row, app, it.scheme_raw, "excel", "OK", f"Status -> {DONE_STATUS_TEXT}")
    return True


def inspect(page, app, delay):
    rows, col = search(page, app, delay)
    tbl = table(page)
    print("\nColumns:", col)
    print("Rows   :", [(r.scheme, r.stage) for r in rows])
    print("\n--- THEAD ---\n", tbl.locator("thead").inner_html()[:1500])
    print("\n--- FIRST ROW ---\n", tbl.locator("tbody tr").first.evaluate("e=>e.outerHTML")[:3000])
    print("\n--- SEARCH INPUTS ---")
    for el in page.locator("input").all():
        print(el.evaluate("e=>e.outerHTML")[:200])


def main():
    ap = argparse.ArgumentParser(description="PM SVANidhi Vendor Scheme Tracker automation")
    ap.add_argument("--excel", default="data.xlsx.xlsx", help="Path to Excel file (default: data.xlsx.xlsx)")
    ap.add_argument("--sheet", default="DFS", help="Worksheet name (default: DFS)")
    ap.add_argument("--cdp", default="http://127.0.0.1:9222", help="Chrome remote-debugging URL")
    ap.add_argument("--dry-run", action="store_true", help="search and report only; change nothing")
    ap.add_argument("--limit", type=int, help="process at most N applications")
    ap.add_argument("--app", action="append", help="only this Application No. (repeatable)")
    ap.add_argument("--delay", type=float, default=0.0, help="seconds to pause after each action (default: 0.0)")
    ap.add_argument("--save-every", type=int, default=5,
                    help="save Excel workbook every N completed applications (default: 5; always saves on exit/error)")
    ap.add_argument("--allow-extra-portal", action="store_true", default=True,
                    help="ignore portal schemes that are not in Excel (enabled by default)")
    ap.add_argument("--strict-portal", action="store_true",
                    help="fail if the portal has extra schemes not listed in Excel")
    ap.add_argument("--no-prefix-check", action="store_true", help="disable Reference No prefix sanity check")
    ap.add_argument("--stop-on-error", action="store_true")
    ap.add_argument("--inspect", metavar="APP_NO", help="print page structure for selector calibration")
    ap.add_argument("--logdir", default="logs")
    args = ap.parse_args()

    target_path = Path(args.excel)
    if not target_path.exists() and target_path.name == "data.xlsx" and Path("data.xlsx.xlsx").exists():
        target_path = Path("data.xlsx.xlsx")

    wb, ws, cols, items, xlsx = load_excel(target_path, args.sheet)

    if not args.dry_run and not args.inspect:
        try:
            open(xlsx, "r+b").close()
        except PermissionError:
            sys.exit(f"Excel file '{xlsx.name}' is open or locked. Please close it in Excel and run again.")
        backup = xlsx.with_name(f"{xlsx.stem}.backup-{datetime.now():%Y%m%d-%H%M%S}{xlsx.suffix}")
        shutil.copy2(xlsx, backup)
        print(f"Backup saved: {backup}")

    groups: "OrderedDict[str, list[Item]]" = OrderedDict()
    for it in items:
        groups.setdefault(it.app, []).append(it)
    todo = [a for a, its in groups.items() if any(i.status == PENDING_STATUS for i in its)]
    if args.app:
        wanted = {appkey(a) for a in args.app}
        todo = [a for a in todo if a in wanted]
    if args.limit:
        todo = todo[: args.limit]
    print(f"{len(todo)} application(s) to process "
          f"({sum(1 for a in todo for i in groups[a] if i.status == PENDING_STATUS)} scheme rows)")

    log = Logger(Path(args.logdir))
    ok, failed = [], []
    unsaved_changes = 0
    save_every = max(1, args.save_every)

    try:
        with sync_playwright() as p:
            try:
                browser = p.chromium.connect_over_cdp(args.cdp)
            except Exception as e:
                sys.exit(f"Cannot connect to Chrome at {args.cdp}. Start Chrome with --remote-debugging-port=9222.\n{e}")
            pages = [pg for c in browser.contexts for pg in c.pages]
            page = next((pg for pg in pages if any(k in pg.url.lower() for k in ["schemestracker", "scheme-tracker", "vendor-scheme", "tracker"])), None)
            if page is None:
                print("ERROR: Could not find the 'Vendor Scheme Tracker' page tab in Chrome.")
                print("Currently open tabs:")
                for i, pg in enumerate(pages, 1):
                    try:
                        print(f"  [{i}] {pg.title()} -> {pg.url}")
                    except Exception:
                        print(f"  [{i}] {pg.url}")
                sys.exit("\nPlease log in and navigate to the 'Vendor Scheme Tracker' page in that Chrome window first.")
            page.bring_to_front()
            disable_animations(page)

            if args.inspect:
                inspect(page, appkey(args.inspect), args.delay)
                return

            for n, app in enumerate(todo, 1):
                vendor = groups[app][0].vendor
                print(f"\n[{n}/{len(todo)}] {app}  {vendor}")
                try:
                    dirty = process_app(page, app, groups[app], args, ws, cols, log)
                    ok.append(app)
                    if dirty:
                        unsaved_changes += 1
                        if unsaved_changes >= save_every:
                            wb.save(xlsx)
                            unsaved_changes = 0
                except PermissionError:
                    sys.exit(f"Could not save Excel file '{xlsx.name}' (is it open in Excel?). Stopped; re-run is safe.")
                except Exception as e:
                    failed.append((app, str(e)))
                    log("", app, "", "app", "ERROR", str(e))
                    try:
                        page.screenshot(path=str(log.dir / f"error-{app}-{datetime.now():%H%M%S}.png"))
                        page.keyboard.press("Escape")
                    except Exception:
                        pass
                    if args.stop_on_error:
                        break
    finally:
        if unsaved_changes > 0 and not args.dry_run and not args.inspect:
            try:
                wb.save(xlsx)
                print(f"\n[SAVED] Flushed final Excel updates to '{xlsx.name}'.")
            except Exception as e:
                print(f"\n[ERROR] Could not save final Excel updates to '{xlsx.name}': {e}")

    print("\n================ SUMMARY ================")
    print(f"Completed : {len(ok)}")
    print(f"Failed    : {len(failed)}")
    for a, msg in failed:
        print(f"  - {a}: {msg}")
    print(f"Log file  : {log.path}")


if __name__ == "__main__":
    main()
