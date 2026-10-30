import openpyxl

def create_sample_files():
    # 1. Create data_with_errors.xlsx (for testing validation logic)
    wb_err = openpyxl.Workbook()
    ws_err = wb_err.active
    ws_err.title = "Sheet1"
    
    headers = ["Name", "Reference", "Status"]
    ws_err.append(headers)

    # Note: openpyxl cell string assignment
    rows_err = [
        ("Ramesh Patel", "100023456789", ""),               # Valid
        ("Priya Sharma", "007894561230", ""),               # Valid (leading zeros & 12 digits)
        ("  Amit   Kumar  ", " 001234567890 ", ""),         # Extra spaces (leading/trailing/double)
        ("Sunita Devi", "100023456789", ""),               # Duplicate Reference of Row 2!
        ("Vikram Singh", "", ""),                           # Missing Reference!
        ("", "998877665544", ""),                           # Missing Name!
    ]

    for r_idx, row in enumerate(rows_err, start=2):
        for c_idx, val in enumerate(row, start=1):
            cell = ws_err.cell(row=r_idx, column=c_idx)
            cell.number_format = '@' # Text format
            cell.value = str(val) if val != "" else ""

    wb_err.save("data_with_errors.xlsx")
    print("Created 'data_with_errors.xlsx' with deliberate validation edge cases.")

    # 2. Create clean data.xlsx (5 valid rows with leading zeros and long numbers)
    wb_clean = openpyxl.Workbook()
    ws_clean = wb_clean.active
    ws_clean.title = "Sheet1"
    ws_clean.append(headers)

    rows_clean = [
        ("Ramesh Patel", "100023456789", ""),
        ("Priya Sharma", "007894561230", ""),               # Leading zero preserved
        ("Amit Kumar", "001234567890", ""),                 # Leading zero preserved
        ("Sunita Devi", "987654321098", ""),
        ("Vikram Singh", "554433221100", "")
    ]

    for r_idx, row in enumerate(rows_clean, start=2):
        for c_idx, val in enumerate(row, start=1):
            cell = ws_clean.cell(row=r_idx, column=c_idx)
            cell.number_format = '@' # Text format
            cell.value = str(val)

    wb_clean.save("data.xlsx")
    print("Created clean 'data.xlsx' with 5 valid test rows.")

if __name__ == "__main__":
    create_sample_files()
